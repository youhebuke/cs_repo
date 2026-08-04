#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build an EDITABLE Kimi K3 deck (v3) that follows the uploaded reference PPT's
6-section structure ("第二讲"), but: native editable text/tables, more figures,
STRENGTHENED pre-training Infra, fewer formulas, detailed per-slide speaker notes.
Run: python3 build_pptx_v3.py  ->  kimi-k3-tech-share-v3.pptx
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")

NAVY = RGBColor(0x1E, 0x3A, 0x5F); NAVY_D = RGBColor(0x0F, 0x24, 0x40)
TEAL = RGBColor(0x0D, 0x94, 0x88); AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_BR = RGBColor(0xF5, 0x9E, 0x0B); VIOLET = RGBColor(0x7C, 0x3A, 0xED)
INK = RGBColor(0x0F, 0x17, 0x2A); GREY = RGBColor(0x64, 0x74, 0x8B)
LGREY = RGBColor(0xE2, 0xE8, 0xF0); PANEL = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); GREEN = RGBColor(0x16, 0xA3, 0x4A)
BG = RGBColor(0xFB, 0xFC, 0xFE)
FONT = "Microsoft YaHei"

SW, SH = 13.333, 7.5
prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]
_idx = [0]


def set_ea(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn("a:" + tag))
        if el is None:
            el = rPr.makeelement(qn("a:" + tag), {}); rPr.append(el)
        el.set("typeface", name)


def bg(slide, color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _set_alpha(shape, pct):
    solid = shape.fill._xPr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int((100 - pct) * 1000))}))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Pt(2))
    return tb, tf


def add_para(tf, text, size=18, color=INK, bold=False, first=False,
             align=PP_ALIGN.LEFT, space_after=6, name=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after)
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    set_ea(run, name)
    return p, run


def fit(path, mw, mh):
    im = Image.open(path); w, h = im.size; r = min(mw / w, mh / h)
    return w * r, h * r


def place_image(slide, path, cx, cy, mw, mh):
    fw, fh = fit(path, mw, mh)
    slide.shapes.add_picture(path, Inches(cx + (mw - fw) / 2),
                             Inches(cy + (mh - fh) / 2), Inches(fw), Inches(fh))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def bullets(tf, items, base=16):
    for i, (text, lvl) in enumerate(items):
        color = INK if lvl == 0 else GREY
        size = base if lvl == 0 else base - 2
        marker = {0: "●  ", 1: "–  ", 2: "·  "}[lvl]
        indent = {0: 0, 1: 0.28, 2: 0.56}[lvl]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5 if lvl == 0 else 3)
        rm = p.add_run(); rm.text = marker
        rm.font.size = Pt(size); rm.font.bold = True
        rm.font.color.rgb = TEAL if lvl == 0 else GREY; set_ea(rm)
        rr = p.add_run(); rr.text = text
        rr.font.size = Pt(size); rr.font.bold = (lvl == 0); rr.font.color.rgb = color
        set_ea(rr)
        pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(int(Inches(indent)))); pPr.set("indent", "0")


def add_table(slide, data, x, y, w, h, fs=13, first_col_accent=True):
    rows, cols = len(data), len(data[0])
    gtbl = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gtbl.table
    # disable built-in banded styling so our per-cell fills show
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Pt(7); cell.margin_right = Pt(7)
            cell.margin_top = Pt(3); cell.margin_bottom = Pt(3)
            tfc = cell.text_frame; tfc.word_wrap = True
            p = tfc.paragraphs[0]
            run = p.add_run(); run.text = str(data[r][c])
            run.font.size = Pt(fs); set_ea(run)
            if r == 0:
                run.font.bold = True; run.font.color.rgb = WHITE
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            else:
                run.font.color.rgb = INK
                cell.fill.solid()
                cell.fill.fore_color.rgb = PANEL if r % 2 == 0 else WHITE
                if c == 0 and first_col_accent:
                    run.font.bold = True; run.font.color.rgb = NAVY
    return tbl


def title_bar(slide, title, sub=None):
    rect(slide, 0.0, 0.0, 0.22, 1.42, TEAL)
    _, tf = textbox(slide, 0.55, 0.26, 12.3, 0.9)
    add_para(tf, title, size=26, color=NAVY, bold=True, first=True)
    rect(slide, 0.58, 1.12, 3.1, 0.045, AMBER_BR)
    if sub:
        _, tf2 = textbox(slide, 0.58, 1.22, 12.0, 0.34)
        add_para(tf2, sub, size=12.5, color=GREY, first=True)


def footer(slide, idx):
    _, tf = textbox(slide, 0.55, 7.08, 11.0, 0.35)
    add_para(tf, "Kimi K3 深度解读 · 结构为什么 / 预训练 / 预训练 Infra", size=9, color=GREY, first=True)
    _, tf2 = textbox(slide, 12.4, 7.08, 0.7, 0.35)
    add_para(tf2, str(idx), size=10, color=GREY, first=True, align=PP_ALIGN.RIGHT)


def takeaway_bar(slide, text):
    rect(slide, 0.6, 6.5, 0.12, 0.5, AMBER_BR)
    rect(slide, 0.72, 6.5, 11.98, 0.5, PANEL)
    _, tf = textbox(slide, 0.9, 6.52, 11.7, 0.46, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, "要点：" + text, size=13, color=NAVY, bold=True, first=True)


def hero_full(slide, dim=42):
    hero = os.path.join(IMG, "cover_hero.png")
    im = Image.open(hero); w, h = im.size; r = max(SW / w, SH / h); pw, ph = w * r, h * r
    slide.shapes.add_picture(hero, Inches((SW - pw) / 2), Inches((SH - ph) / 2), Inches(pw), Inches(ph))
    ov = rect(slide, 0, 0, SW, SH, NAVY_D); _set_alpha(ov, dim)


# ---------------- slide constructors ----------------
def slide_title():
    s = prs.slides.add_slide(BLANK); bg(s, NAVY_D); hero_full(s, dim=30)
    _, tf = textbox(s, 0.75, 1.55, 11.8, 2.4)
    add_para(tf, "Kimi K3 深度解读", size=46, color=WHITE, bold=True, first=True, space_after=6)
    add_para(tf, "为什么这样设计 · 预训练 · 预训练 Infrastructure", size=22,
             color=RGBColor(0x93, 0xC5, 0xFD), bold=True, space_after=6)
    add_para(tf, "Kimi K3: Open Frontier Intelligence（2026.07）技术报告精读 · 第二讲", size=13,
             color=RGBColor(0xCB, 0xD5, 0xE1))
    chips = [("2.8T", "总参数"), ("104B", "激活参数"), ("1M", "上下文"),
             ("896选16", "MoE 专家"), ("2.5×", "扩展效率 vs K2")]
    x = 0.8
    for big, small in chips:
        rect(s, x, 4.35, 2.15, 1.15, NAVY)
        _, tf2 = textbox(s, x, 4.45, 2.15, 1.0, anchor=MSO_ANCHOR.MIDDLE)
        add_para(tf2, big, size=22, color=AMBER_BR, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=2)
        add_para(tf2, small, size=11, color=WHITE, align=PP_ALIGN.CENTER)
        x += 2.35
    _, tf3 = textbox(s, 0.8, 5.85, 11.6, 0.5)
    add_para(tf3, "本讲侧重：结构设计的动机 / 预训练与其系统实现 · 结构基础与推理服务从略", size=13,
             color=RGBColor(0xE2, 0xE8, 0xF0), first=True)
    _idx[0] = 1; footer(s, 1)
    notes(s, "【开场 · 约1分钟】这是 Kimi K3 系列分享的第二讲。第一讲已经把结构『是什么』过了一遍，"
             "所以今天把结构基础压缩，重点讲三件事：每个结构选择『为什么这样设计』背后的痛点和取舍；"
             "预训练的数据与方法论；以及支撑 2.8T 训练的预训练 Infrastructure——这是本讲最重的部分。"
             "强化学习简化过，推理/在线服务这次不讲。\n\n"
             "先给一个尺度感：2.8T 总参、104B 激活（只占 3.7%，非常稀疏）、1M 上下文、896 选 16 的 MoE、"
             "相对 K2 约 2.5× 的整体扩展效率——这些数字后面都会展开『为什么能做到』。"
             "提醒：每页我都写了备注，可当逐字稿；含公式的地方备注里有逐符号解释。")


def slide_section(num, title, sub, note):
    s = prs.slides.add_slide(BLANK); bg(s, NAVY_D); hero_full(s, dim=48)
    _, tf0 = textbox(s, 1.05, 2.5, 3, 1.2)
    add_para(tf0, num, size=64, color=RGBColor(0x5E, 0xEA, 0xD4), bold=True, first=True)
    rect(s, 0.7, 3.9, 0.16, 1.5, AMBER_BR)
    _, tf = textbox(s, 1.05, 3.85, 11.0, 1.8)
    add_para(tf, title, size=38, color=WHITE, bold=True, first=True, space_after=8)
    add_para(tf, sub, size=16, color=RGBColor(0xCB, 0xD5, 0xE1))
    _idx[0] += 1; notes(s, note)
    return s


def content(title, sub, blts=None, image=None, layout="right", table=None,
            col_h=None, takeaway=None, note=""):
    s = prs.slides.add_slide(BLANK); bg(s)
    title_bar(s, title, sub)
    ty = 1.62
    if layout == "table":
        rows = len(table)
        th = min(4.6, 0.52 * rows + 0.1)
        add_table(s, table, 0.6, ty, 12.1, th, fs=13.5)
        if blts:
            _, tf = textbox(s, 0.6, ty + th + 0.15, 12.1, 6.4 - (ty + th))
            bullets(tf, blts, base=14)
    elif layout == "table_img":
        rows = len(table)
        th = min(4.4, 0.52 * rows + 0.1)
        add_table(s, table, 0.6, ty, 7.2, th, fs=12.5)
        place_image(s, os.path.join(IMG, image), 7.95, 1.5, 5.0, 5.0)
        if blts:
            _, tf = textbox(s, 0.6, ty + th + 0.15, 7.2, 6.4 - (ty + th))
            bullets(tf, blts, base=13)
    elif layout == "two":
        place_image(s, os.path.join(IMG, image[0]), 0.5, 1.55, 6.1, 3.7)
        place_image(s, os.path.join(IMG, image[1]), 6.75, 1.55, 6.1, 3.7)
        if blts:
            _, tf = textbox(s, 0.6, 5.35, 12.1, 1.1)
            bullets(tf, blts, base=13)
    elif layout == "right" and image:
        _, tf = textbox(s, 0.6, ty, 5.9, 5.0); bullets(tf, blts, base=15.5)
        place_image(s, os.path.join(IMG, image), 6.65, 1.5, 6.35, 5.1)
    elif layout == "left" and image:
        place_image(s, os.path.join(IMG, image), 0.45, 1.5, 6.35, 5.1)
        _, tf = textbox(s, 7.05, ty, 5.85, 5.0); bullets(tf, blts, base=15.5)
    elif layout == "big" and image:
        place_image(s, os.path.join(IMG, image), 0.6, 1.4, 12.1, 3.35)
        _, tf = textbox(s, 0.6, 4.9, 12.1, 1.55); bullets(tf, blts, base=13)
    else:
        _, tf = textbox(s, 0.6, ty, 12.2, 5.0); bullets(tf, blts, base=17)
    if takeaway:
        takeaway_bar(s, takeaway)
    _idx[0] += 1; footer(s, _idx[0]); notes(s, note)
    return s


# =====================================================================
def build():
    slide_title()

    # roadmap
    content("今天的重点（第二讲）", "六个部分，结构『为什么』/ 预训练 / 预训练 Infra 是重心",
            table=[["#", "章节", "内容", "配时"],
                   ["01", "设计哲学与决策总表", "三维信息流；一页看全『选择→原因』", "5 min"],
                   ["02", "结构设计的原因（精简）", "序列/深度/宽度：为什么，而非是什么", "12 min"],
                   ["03", "预训练（扩充）", "数据 / scaling law 方法论 / 1M 课程 / 稳定性", "12 min"],
                   ["04", "预训练 Infra（重点）", "FlashKDA / KCP / MoonEP / 显存 / ViT", "15 min"],
                   ["05", "强化学习（简化）", "三阶段 / 环境 / MOPD 一带而过", "7 min"],
                   ["06", "评测与总结", "成本效率 / 案例 / 方法论启示", "3 min"]],
            layout="table",
            blts=[("相较上一讲：结构介绍与推理服务收缩，预训练与预训练 Infra 显著扩充", 0)],
            note="【路线图 · 约1分钟】六个部分，配时在表里。结构的『为什么』12 分钟、预训练 12 分钟、"
                 "预训练 Infra 15 分钟——这三块是重心，合计接近 40 分钟。强化学习 7 分钟快速过，评测总结 3 分钟。\n\n"
                 "和上一讲的差异：结构基础和推理服务这次收缩，把时间让给预训练与它的系统实现。"
                 "控场提示：如果前面结构部分讨论热烈，可以从 02 节里标了『详见备注』的页快速跳过，保证 03、04 两节的时间。")

    # ============ 01 ============
    slide_section("01", "设计哲学与决策总表",
                  "先建立统一视角：K3 的每个改动都在扩展某个维度的『信息流』",
                  "【过场 · 约30秒】这一节只做一件事：给一个统一的观察框架，让后面所有结构选择都能挂靠上去。"
                  "核心观点——K3 不是一堆零散 trick，而是围绕『信息在三个维度上如何流动』组织的：序列（token 之间）、"
                  "深度（层之间）、宽度（通道/专家之间），每个创新对应某个维度的一个瓶颈。然后用一页『决策总表』当整场地图。")

    content("统一视角：在三个维度上扩展信息流", "序列 / 深度 / 宽度，各有痛点与方案",
            table=[["维度", "瓶颈", "K3 方案"],
                   ["序列 token mixing", "1M 下 softmax 计算 O(T²) 与 KV cache O(T) 爆炸", "3×KDA + 1×Gated MLA"],
                   ["深度 layer mixing", "残差把历史层等权压成单态，深层看不清早期", "Attention Residuals"],
                   ["宽度 channel mixing", "扩专家池 → 通信/带宽/稳定性恶化", "Stable LatentMoE"]],
            image="f_threeaxis.png", layout="table_img",
            blts=[("每 block：3×(KDA+MoE) + 1×(Gated MLA+MoE)，末层强制全局", 0),
                  ("本讲对『是什么』点到为止，重点在后面『为什么』", 0)],
            takeaway="KDA 跨 token · AttnRes 跨 layer · LatentMoE 跨 channel。",
            note="【三维框架 · 约1.5分钟 · 全场地图】右图是架构总图，给一个空间锚点，不逐一讲模块（上一讲讲过）。"
                 "用左表把三维框架讲清楚：\n"
                 "序列维——1M 下标准 softmax 计算是 O(T²)、KV cache 是 O(T)，都撑不住，K3 用 3 层 KDA（线性、固定状态）配 1 层全局 MLA。\n"
                 "深度维——标准残差把之前所有层等权加起来压成一个状态，深层看不清早期信息，K3 用 Attention Residuals，让每层对前面各层做注意力。\n"
                 "宽度维——扩专家池会让通信/带宽/稳定性恶化，K3 用 Stable LatentMoE 在压缩空间里路由。\n\n"
                 "让听众记住这张表——后面每个结构选择都挂在这三根轴上。")

    content("K2 → K3：为什么值得从头重调", "加深 + 加稀疏，而非加宽",
            table=[["", "Kimi K2", "Kimi K3", "Δ"],
                   ["层 / 总参 / 激活", "61 / 1.04T / 32.6B", "93 / 2.78T / 104.2B", "↑52 / 167 / 220%"],
                   ["注意力", "61× MLA", "69× KDA + 24× MLA", "混合线性"],
                   ["MoE", "384 选 8 + 1", "896 选 16 + 2 (latent 3584)", "稀疏度 56"],
                   ["激活 / 上下文", "SwiGLU / 128K", "SiTU-GLU / 1M", "8× ctx"]],
            image="f_scaling.png", layout="table_img",
            blts=[("hidden 维不变(7168)，靠加深+加稀疏扩张；激活仅占 3.7%", 0),
                  ("2.5× 是『整体』效率（架构+数据+超参共同贡献），OOD 验证集拟合", 0)],
            takeaway="改动大到改变了『最优训练 regime』→ 必须重做 scaling law 搜索。",
            note="【K2→K3 · 约1.5分钟】回答一个高层问题：K3 相对 K2 强在哪、为什么改动这么大。"
                 "左表看规模：层数 +52%、总参 +167%、激活 +220%，但 hidden 维度没变（7168）——扩张主要靠『加深』和『加稀疏』，"
                 "激活只占总参 3.7%，非常稀疏。右图是关键证据：K2/K3 各自在 held-out、OOD 验证集上拟合 scaling law，"
                 "同样 loss，K3 需要的算力约是 K2 的 1/2.5——这就是『2.5× 整体扩展效率』的来源。"
                 "强调两点：2.5× 是架构+数据 recipe+训练超参共同贡献，不是单点；且改动大到改变了最优训练配置，所以必须重做 scaling law。")

    content("一页看全：每个选择背后的原因", "整场索引 · 痛点 → 选择 → 为什么",
            table=[["维度 / 模块", "痛点", "选择", "为什么是它"],
                   ["序列 token mixing", "1M 下 softmax 太贵", "3 KDA : 1 MLA 混合", "线性扛长程 + 全局保精确检索"],
                   ["KDA 衰减门", "chunkwise 对角块有慢路径", "lower-bounded decay", "数值压进 BF16 → 全块走 Tensor Core"],
                   ["全局层位置编码", "扩上下文要重调 RoPE/YaRN", "NoPE", "位置由 KDA 隐式携带 → 零调参外推"],
                   ["深度 layer mixing", "残差稀释早期层", "Attention Residuals", "深度维照搬序列维的 attention 革命"],
                   ["宽度 MoE", "扩专家→通信/带宽爆炸", "Stable LatentMoE", "latent 压缩 + 归一/有界/分位数均衡"],
                   ["优化器", "大规模各头更新失衡", "Per-Head Muon", "逐头正交化均衡更新尺度"]],
            layout="table",
            note="【决策总表 · 约2分钟 · 本节地图】这页是整场索引，把主要设计选择、各自解决的痛点、以及『为什么是这个选择』"
                 "一次性列出来。不用逐行念，挑两条点一下：\n"
                 "序列 token mixing——痛点是 1M 下 softmax 太贵，选混合注意力，原因是『线性层扛长程、全局层保精确检索』。\n"
                 "KDA 衰减门——痛点是 chunkwise 里对角块有一条慢路径，选 lower-bounded decay，原因很硬核：把数值范围压进 BF16，"
                 "让所有块都能走 Tensor Core。这是『算法为硬件让路』的典型，后面 Infra 的 FlashKDA 会呼应。其余在 02 节展开。")

    # ============ 02 ============
    slide_section("02", "结构设计的原因（精简）",
                  "不讲『是什么』，只讲『为什么这样、为什么不是别的』",
                  "【过场 · 约30秒】结构部分刻意精简『是什么』。每一页用『痛点 → 选择 → 为什么不是别的方案』来讲动机。"
                  "公式大幅压缩，只在 KDA 那页保留一个招牌公式，其余用图；需要展开数学的地方我把逐符号解释写进了备注。"
                  "这节 12 分钟、8 页，平均一页一分半，注意别恋战。")

    content("序列①：为什么是『混合』，而不是纯线性/纯全局", "把成本与检索解耦",
            image="f_hybrid_tradeoff.png", layout="left",
            blts=[("纯全局 softmax：精确但最贵（O(T²)、KV O(T)），1M 下不可行", 0),
                  ("纯线性：便宜但有损——固定状态是压缩，精确检索会退化", 0),
                  ("K3：3 KDA : 1 MLA，多数层便宜、每 4 层一个全局『锚点』", 0),
                  ("只有 1/4 的层贡献随 T 增长的 cache（还是低秩 MLA）", 1),
                  ("结论：混合把『成本』与『检索』解耦，逼近左上角理想区", 0)],
            takeaway="不是一条道走到黑，而是让不同层各司其职。",
            note="【序列① · 约1.5分钟】序列维第一个『为什么』：为什么混合。看左图这张二维图：横轴长序列成本（越右越贵），"
                 "纵轴精确检索能力（越上越强）。右上角是全 softmax——检索最强但最贵，1M 下算不起也存不起；"
                 "左下角是纯线性（KDA/GDN）——状态固定、非常便宜，但有损压缩，做大海捞针式精确检索会退化。"
                 "K3 选在中间偏左上：3 层 KDA 配 1 层 Gated MLA。直觉是——大部分 token 混合交给便宜的线性层，"
                 "每 4 层插一个全局注意力当『锚点』，周期性重看完整上下文。系统收益很直接：只有 1/4 的层贡献随长度增长的 cache，"
                 "而且那还是低秩的 MLA。")

    content("序列②：KDA —— 唯一需要记住的一个公式", "用固定状态换掉增长的 KV cache",
            image="f_kda_state.png", layout="right",
            blts=[("把『不断增长的 KV cache』换成『固定大小状态 S』——长序列便宜的关键", 0),
                  ("delta rule：写入前先擦掉该 key 方向的旧值", 0),
                  ("= 对在线回归目标走一步梯度下降", 1),
                  ("通道级遗忘门 Diag(α)：每个 key 通道独立决定遗忘多少", 0),
                  ("比 GDN 的标量门更细；配 ShortConv + L2Norm + 满秩输出门", 1)],
            takeaway="固定大小状态——记住这点，Infra 的 KCP 与前缀缓存都建立在它之上。",
            note="【序列② · 约2分钟 · 全场唯一重点公式，务必讲透】上图对比两种记忆方式：左边 softmax 把每个历史 token 的 k、v 都存下来，"
                 "cache 随序列线性增长；右边 KDA 只维护一个固定大小的状态矩阵 S，边读边更新——这是长序列便宜的关键。\n\n"
                 "公式（逐符号，看备注慢讲）：S_t = (I − β_t·k_t·k_tᵀ)·Diag(α_t)·S_{t−1} + β_t·k_t·v_tᵀ，输出 õ_t = S_tᵀ·q_t。\n"
                 "S_t 是 t 时刻的递归状态矩阵，一份持续更新的记忆摘要，大小固定、与序列长度无关；"
                 "(I − β_t·k_t·k_tᵀ) 是 delta rule 的擦除项——写入前先沿 key 方向删掉旧关联，等价于对在线回归目标走一步梯度下降；"
                 "Diag(α_t) 是通道级遗忘门，每个 key 通道独立决定遗忘多少，比 GDN/Mamba-2 的标量门精细得多。"
                 "请记住『固定大小状态』——Infra 的 KCP 和前缀缓存都建立在它之上。")

    content("序列③：一个『为硬件让路』的小改动", "改门控值域，删掉一整条 kernel 慢路径",
            image="f_kda_decay.png", layout="right",
            blts=[("问题：chunkwise 要除以累积衰减 1/Γ，α 连乘 → 可无界增长、溢出", 0),
                  ("Kimi Linear 老办法：对角块逐位置对算 —— 一条慢路径，上不了 Tensor Core", 0),
                  ("K3：把衰减门下界卡在 gmin = −5（左图红线）", 0),
                  ("16-token 块累积衰减 ∈ (−80,0) → 倒数 < e⁸⁰ 落进 BF16", 1),
                  ("收益：对角块也能走稠密 Tensor Core，慢路径被删除", 0)],
            takeaway="改一个门控函数的值域，换来 kernel 一整条慢路径被删除（详见 Infra·FlashKDA）。",
            note="【序列③ · 约1.5分钟】一个很能体现『算法-系统协同』的小改动，用图讲。背景：KDA 训练用 chunkwise 形式，"
                 "块内并行、块间递归。块内计算里 key 要除以累积衰减 Γ，而 Γ 是很多 α∈(0,1) 连乘的结果，连乘趋近 0，"
                 "倒数 1/Γ 可能无限变大、有限精度下溢出。Kimi Linear 老办法是切成 16-token 小 tile，非对角 tile 能走 Tensor Core，"
                 "但对角 tile 还得逐位置对地算——kernel 里一条慢路径。\n\n"
                 "K3 的改法（左图）：把衰减门从『无下界的负 Softplus』换成『有下界的缩放 sigmoid』，下界卡在 gmin=−5。"
                 "于是 16-token 块累积衰减落在 (−80,0)，倒数小于 e⁸⁰，落进 BF16 范围。结果（右图）：对角块也能走稠密 Tensor Core，"
                 "那条慢路径被彻底删掉。这就是『算法为硬件让路』——这条到 Infra 的 FlashKDA 会再展开。")

    content("序列④：为什么敢用 NoPE + 给全局层加门", "职责分离：KDA 扛长程+位置，MLA 做全局检索",
            layout="none",
            blts=[("Gated MLA：全局注意力层，沿用 DeepSeek-V2 的 MLA（KV 压成低维 latent 再重建）", 0),
                  ("KV cache 小；K3 再加一个输入依赖的满秩 sigmoid 输出门（与 KDA 一致）", 1),
                  ("NoPE：MLA 层完全不加位置编码", 0),
                  ("为什么敢：位置信息由中间的 KDA 层通过递归 + 衰减门隐式携带", 1),
                  ("红利：上下文从 8K 扩到 1M 全程零位置编码调参（无 RoPE rescale / YaRN）", 1),
                  ("取舍主线：KDA 便宜地扛长程 + 位置感，MLA 周期性做无位置约束的全局检索", 0)],
            takeaway="NoPE 是 K3 能『直接外推到 1M』的关键之一。",
            note="【序列④ · 约1.5分钟】两个『为什么』，不放公式。第一，全局层为什么用 Gated MLA：MLA 是 DeepSeek-V2 提出的，"
                 "把每个 token 的 K、V 压成低维 latent 缓存、用时上投影重建，KV cache 小很多同时保留全局注意力；"
                 "K3 在它基础上加一个输入依赖的满秩 sigmoid 输出门，和 KDA 保持一致，让每个 token 调制自己从全局注意力读到的通道。\n\n"
                 "第二，更有意思——为什么敢用 NoPE（全局层完全不加位置编码）。常识里注意力必须有位置编码，否则分不清顺序；"
                 "但 K3 的位置信息已经由中间的 KDA 层通过递归和衰减门隐式携带了，MLA 只负责无位置约束的全局内容检索。"
                 "红利很大：上下文从 8K 扩到 1M 全程不用碰位置编码——不用 RoPE rescale、不用 YaRN 插值，直接外推。"
                 "这页的主线是『职责分离』。")

    content("深度：残差是『深度方向的 RNN』，那就换成 attention", "把序列维的 attention 革命搬到深度维",
            image="f_attnres.png", layout="left",
            blts=[("痛点：标准残差把之前所有层等权累加成一个状态 → 深层信息被稀释", 0),
                  ("PreNorm 下幅值还随深度无界增长，破坏稳定", 1),
                  ("类比：这正是 RNN 沿时间的瓶颈；序列维早已用 attention 解决 → 深度维照做", 0),
                  ("做法：每层用一个可学习 pseudo-query，对前面各层输出做 softmax 注意力", 0),
                  ("Block 版：分块降开销 O(Ld) → O(Nd)，N≈8 恢复大部分收益，推理开销 <2%", 0)],
            takeaway="标准残差=深度维 RNN；AttnRes=深度维 softmax。",
            note="【深度 · 约1.5分钟】用图讲，公式放备注。痛点（图 a）：标准残差 h_{l+1}=h_l+f_l(h_l)，"
                 "等于把之前所有层的输出等权加起来压进一个不断累积的状态，层一深早期信息就被稀释，PreNorm 下幅值还无界增长——"
                 "这本质就是 RNN 沿时间维的瓶颈。关键类比：序列维当年用 attention 替换了 RNN 递归，让每个位置选择性看全部历史，"
                 "K3 把同一招用到深度维，就是 Attention Residuals（图 b）。做法：每层有一个可学习 pseudo-query，"
                 "对前面各层输出做 softmax 注意力，选择性聚合。Block 版（图 c）把层分块，只在块间做注意力，"
                 "开销从 O(Ld) 降到 O(Nd)，N≈8 就能恢复大部分收益，推理延迟只多不到 2%。")

    content("宽度①：为什么 896 个专家还能训得起", "解耦『模型全宽』与『专家宽度』",
            image="f_latentmoe.png", layout="left",
            blts=[("痛点：常规 MoE 里每个被选专家都吃完整 d 维 → 通信/权重带宽随激活数线性涨", 0),
                  ("LatentMoE：共享专家保留全宽通路；路由专家在 ℓ=d/2=3584 的 latent 空间工作", 0),
                  ("dispatch/combine 通信减半、权重带宽减半；激活专家 8→16、专家 384→896", 1),
                  ("极端稀疏 + 2.8T 放大两个失效模式", 0),
                  ("① W↓→门控FFN→W↑ 近 4 连乘的病态链 → 内部激活爆炸", 1),
                  ("② 近千专家的负载均衡超出传统 sign 偏置的可控范围", 1)],
            takeaway="latent 压缩让『更多、更小、更专』的专家变得可负担。",
            note="【宽度① · 约1.5分钟】为什么能把专家堆到 896 还训得起。痛点：常规 MoE（DeepSeekMoE 那种）里每个被选专家都要接收完整 d 维表示，"
                 "每多激活一个专家，all-to-all 通信量和读专家权重的带宽都线性增加，限制了往『更多、更小、更专』方向扩。\n\n"
                 "LatentMoE 的解法：解耦『模型全宽 d』和『路由专家工作宽度 ℓ』。共享专家（每个 token 都过）保留全宽 d 通路负责共性变换；"
                 "路由专家在压缩的 latent 空间工作，K3 取 ℓ=d/2=3584。于是 dispatch/combine 通信减半、权重带宽减半，"
                 "才敢把激活专家从 8 提到 16、专家池从 384 扩到 896。但极端稀疏 + 2.8T 放大两个失效模式，下一页两个稳定化设计来治。")

    content("宽度②：两个稳定化设计——SiTU-GLU 与 Quantile Balancing", "一个压激活幅值，一个压负载失衡",
            image=["f_situ.png", "f_qb.png"], layout="two",
            blts=[("SiTU-GLU：给 SwiGLU 两分支各戴光滑软帽 β·tanh(x/β)，近原点≈SwiGLU、输出≤β₁β₂=100 → 为 MXFP4/8 兜底", 0),
                  ("Quantile Balancing：把均衡看作最优指派，专家偏置有闭式解=目标负载分位数；无学习率超参、几步收敛、直方图全局估计", 0)],
            takeaway="都服务于『极端稀疏下的稳定训练』。",
            note="【宽度② · 约2分钟 · 两个都不放公式，讲思路+备注给公式】左边 SiTU-GLU，解决激活爆炸。SwiGLU 是 gate·value 两分支相乘，"
                 "两个因子都无界，大坐标一共振就产生离群值，低精度（MXFP4/8）下易溢出。SiTU-GLU 给两分支都套光滑软帽 β·tanh(x/β)。"
                 "公式（备注）：SiTU-GLU(x)=[β₁·tanh(Wg x/β₁)⊙σ(Wg x)]⊙[β₂·tanh(Wu x/β₂)]，β₁=4、β₂=25。"
                 "三条性质：近原点一阶等于 SwiGLU（不改原点行为）；输出上界 β₁β₂=100；相比硬 clamp 在非饱和区保留非零梯度。\n\n"
                 "右边 Quantile Balancing，解决近千专家的负载均衡。传统 sign 偏置（负载多减一点、少加一点）步长大震荡、小收敛慢，"
                 "近千专家失控。QB 把均衡看成最优指派问题，对偶推出专家偏置的闭式解正好是 margin 分布的分位数（右图红线），"
                 "无学习率类超参、几步收敛；全局分位数用直方图估计——一次 all-reduce 求 bin 计数，和 token 分片方式无关。")

    content("横向对照：三家长上下文注意力路线", "两条根本路线",
            table=[["", "路线", "代表", "机制要点"],
                   ["线性状态混合", "RNN 式固定状态 + 周期全局层", "Kimi K3 / Qwen3.5", "K3 通道级 Diag(α)+ShortConv+满秩门；Qwen 标量 GDN"],
                   ["压缩+稀疏检索", "保 softmax，压缩 KV + top-k 选择", "DeepSeek-V4", "CSA 4× / HCA 128× 压缩 + 轻量 indexer 稀疏选"]],
            layout="table",
            blts=[("K3 与 Qwen 走线性混合；DeepSeek 走稀疏 softmax —— 无绝对优劣，是两种工程审美", 0)],
            takeaway="线性状态混合 vs 压缩稀疏检索：两条根本路线。",
            note="【横向对照 · 约1.5分钟 · 本节收口】把三家 2026 旗舰的长上下文注意力放一起建立坐标。两条根本路线："
                 "线性状态混合（K3 与 Qwen3.5）——用 RNN 式固定状态承担大部分 token mixing，周期插全局层保精确检索；"
                 "区别在门控粒度，K3 的 KDA 是通道级 Diag(α)、还多了 ShortConv/L2Norm/满秩输出门，Qwen 的 GDN 是标量 α。"
                 "压缩+稀疏检索（DeepSeek-V4）——保留 softmax 语义，但把 KV 先压缩（CSA 4×、HCA 128×），再用轻量 indexer 做 top-k 稀疏选择，"
                 "好处是每层都保留精确检索。没有绝对优劣，两种工程审美，团队可按自己需求选。")

    # ============ 03 ============
    slide_section("03", "预训练",
                  "数据 recipe · scaling law 方法论 · 1M 长上下文课程 · 稳定性",
                  "【过场 · 约30秒】进入预训练，这节我扩充了，因为它是把架构潜力兑现成能力的关键，且有几个方法论值得团队直接借鉴。"
                  "四条主线：数据（四域+视觉、rephrasing、程序化多模态）、scaling law 方法论（尤其 cosine vs WSD 的公平比较）、"
                  "1M 长上下文的四阶段课程与长依赖合成、以及训练稳定性的一组手段。")

    content("预训练数据：不只是『更多』，而是『更会喂』", "语料构成 + 两个数据工程",
            layout="none",
            blts=[("文本四域：Web / Code / Math / Knowledge，各经规则+分类器打分+去重，配比靠小模型消融确定", 0),
                  ("视觉语料：caption / 图文交错 / OCR / 感知 / 视频 / visual coding；坐标给绝对+归一化两种格式", 0),
                  ("Rephrasing（承自 K2）：对知识/数学做风格与视角多样化改写 + 对原文保真校验", 0),
                  ("让同一份知识以不同角度多次曝光，而非简单重复 → 提高每 token 知识密度", 1),
                  ("程序化多模态：代码片段配其渲染结果（SVG/3D/网页/游戏/CAD）", 0),
                  ("→ 支撑『写代码看效果』的 vision-in-the-loop 能力", 1)],
            takeaway="配比靠消融、改写要保真、程序化数据造就视觉编程能力。",
            note="【数据 · 约2分钟】数据侧三件事。第一，语料构成：文本四域各经规则+分类器打分+去重，域间配比靠小模型消融确定（不是拍脑袋）；"
                 "视觉语料包含 caption、图文交错、OCR、感知、视频、visual coding，坐标监督给绝对和归一化两种格式让定位对分辨率鲁棒。\n\n"
                 "第二，两个『更会喂』的数据工程：Rephrasing（承自 K2）——对知识和数学做风格与视角多样化改写、分块自回归生成、"
                 "再和原文做保真校验，目的是让同一份知识以不同角度多次曝光而非简单重复，提高每 token 的知识密度；"
                 "程序化多模态——代码片段配其渲染结果，覆盖 SVG、3D、网页、游戏、CAD，这正是 K3『写代码看渲染效果再改』能力的来源。")

    content("方法论：一个值得抄的 scaling law 实验设计", "cosine vs WSD 的公平比较",
            image="f_scaling.png", layout="right",
            blts=[("改动大到改变最优训练 regime → 重做 scaling law 搜索（batch / LR / TPP / 形状）", 0),
                  ("cosine vs WSD：两种调度的最优超参（峰值 LR、batch）差异很大", 0),
                  ("用同一组超参比较，会不公平地偏向恰好更适配它的那个", 1),
                  ("正确做法：分别给两者各自调到最优，再比 → 结论才是 cosine 更优", 0)],
            takeaway="比较任意训练技术前，先分别各自调到最优——几乎适用于所有 A/B 实验。",
            note="【scaling law 方法论 · 约2分钟 · 本节最想让人带走的方法论，慢讲】背景：K3 架构、数据、训练都变了，最优训练配置也变了，"
                 "所以没沿用 K2 超参，而是重做一轮 scaling law 搜索，重调 batch、峰值 LR、tokens-per-parameter、模型形状。\n\n"
                 "但这页核心不是结论，是一个实验设计的坑。他们比较 cosine decay 和 WSD 两种学习率调度，关键观察是——两种调度各自的最优超参差别很大，"
                 "最优峰值 LR、最优 batch 都不一样。所以如果固定一组超参去比，你比的其实是『超参对谁更适配』，而不是调度本身。"
                 "正确做法是分别把两者各自调到最优再比，结论才是 cosine 最终 loss 更低。这条几乎适用于所有 A/B 实验，强烈建议团队照做。")

    content("1M 上下文：四阶段课程 + 强制长依赖", "课程经济学 + 长≠长程能力",
            image="f_longctx.png", layout="right",
            blts=[("四阶段：预训练 8K→64K，cooldown 256K→1M（把最贵的长序列计算集中在一小段预算）", 0),
                  ("长数据清洗：精确+模糊去重 / 视频帧感知哈希 / 结构校验；真长文稀缺→上采样", 0),
                  ("长 ≠ 长程能力：合成『必须跨全 1M 取证才能解』的拼接任务，强制 attention 在目标尺度训练", 0),
                  ("NoPE + KDA → 全程零位置编码调参（无 RoPE rescale / YaRN）——呼应结构节", 0)],
            takeaway="NoPE 负责『能外推』，合成长依赖负责『真的会用长上下文』。",
            note="【1M 上下文 · 约1.5分钟】专讲 1M 怎么练出来。四阶段课程（右图）：预训练窗口 8K→64K，cooldown 256K→1M，"
                 "核心是『课程经济学』——长序列计算极贵，只在总预算的一小部分做、逐步拉长、渐进适应。"
                 "数据处理：自然长文/长视频有大量低质（近重复、二进制块、截断、无效日志），用专门管线清洗（精确+模糊去重、视频帧感知哈希、结构校验），"
                 "真长文稀缺还要上采样。最重要还是那句『长≠长程能力』：单纯喂长文本模型可能退化成只看局部，"
                 "所以合成『必须跨越整个 1M 取证才能解』的拼接任务，强制 attention 在目标尺度训练。位置上 NoPE+KDA 让全程零位置编码调参，呼应结构节。")

    content("训练稳定性：一组『配套』手段", "稳不是单点，是组合",
            table=[["手段", "作用 / 为什么"],
                   ["Per-Head Muon", "Q/K/V 动量按 head 分块逐头正交化，均衡各头更新尺度"],
                   ["K2 weight clipping", "沿用 K2 权重裁剪，抑制爆炸"],
                   ["Quantile Balancing", "MoE 负载均衡，避免专家训不充分拖慢 EP"],
                   ["MoonViT-V2 从零训练", "不用 SigLIP 初始化：联合训练更稳，视觉评测持平"],
                   ["QAT 从 SFT 起", "训练即感知 MXFP4/8 量化，训推一致"]],
            layout="table",
            blts=[("Per-Head Muon：整块正交化会让梯度大的头主导共享更新方向 → 逐头让各头更新尺度对齐", 0),
                  ("MoonViT-V2 结论很有信息量：规模够大时，对比学习预训练不是多模态 LLM 的必要初始化", 0)],
            takeaway="每一项都在回答『大规模下怎么不炸』。",
            note="【稳定性 · 约2分钟】稳定性不是单点，是一组配套手段。Per-Head Muon（重点，也算这页的公式替代讲解）："
                 "Muon 核心是对矩阵参数的动量做 Newton–Schulz 正交化再更新，token 效率比 AdamW 高；K3 改进是逐头——"
                 "Q/K/V 投影不当成一整块正交化，而是沿 head 维切开、每头单独做。为什么？整块正交化会把所有头耦合，"
                 "梯度或动量尺度大的头会主导那个共享更新方向、小头归一化不足；逐头让各头更新尺度对齐，大规模更稳，"
                 "顺带 Newton–Schulz 在瘦高小矩阵上更省。\n\n"
                 "K2 的 weight clipping 继续用抑制爆炸；Quantile Balancing 从稳定性角度再点一次——均衡不好会让部分专家训不充分、拖慢 EP。"
                 "MoonViT-V2 从零训练是个反常规但很有价值的决策：不用 SigLIP 对比学习初始化视觉塔，而是从零用 next-token prediction 和 LLM 一起训，"
                 "原因就是稳定性——SigLIP 初始化联合训练时梯度范数持续偏高、频繁尖刺，从零训全程平稳且评测持平。"
                 "结论：规模够大时，对比学习预训练不是多模态 LLM 的必要初始化。")

    # ============ 04 · INFRA (strengthened) ============
    slide_section("04", "预训练 Infrastructure（重点）",
                  "让 2.8T × 1M 训得起、训得稳、训得快 —— 算法与系统的协同设计",
                  "【过场 · 约40秒 · 本讲最重的一节，15 分钟】K3 同时踩中三个平时很少叠加的系统难题：混合 KDA 注意力（全新算子）、"
                  "3T 级稀疏训练、百万 token 序列。组织方式：先给并行组合和三大痛点的总览，然后逐个击破——"
                  "FlashKDA/KCP 解决『新算子怎么高效』；MoonEP 解决『专家负载不均』；显存工程解决『装不下』；ViT 优化解决『多模态计算方差』，"
                  "最后和 DeepSeek 的训练 infra 做对照。这节图多公式少，重点理解『每个系统设计对应哪个痛点、原理是什么』。")

    content("总览：并行组合与三大痛点", "四层并行 + 三个待解痛点",
            image="f_parallel.png", layout="left",
            blts=[("PP + VPP：93 层切流水段，interleaved 1F1B 减 bubble", 0),
                  ("EP：896 专家分片到各 rank，all-to-all 做 dispatch/combine", 0),
                  ("CP（KCP）：1M 序列切段，固定大小 all-gather 同步递归状态（KDA 特有）", 0),
                  ("ZeRO-1 DP + Pipeline ZeRO-2：优化器状态/梯度分片，梯度可 offload CPU", 0),
                  ("三大痛点 = 本节解题顺序：① EP 负载不均 ② 显存超预算 ③ ViT 计算方差", 0)],
            takeaway="这套组合的精髓：切完之后每条通信都被某段计算盖住。",
            note="【Infra 总览 · 约1.5分钟】先建立全局观。K3 预训练叠了四层并行（上半图）：PP+虚拟阶段 VPP 把 93 层切流水段、"
                 "用 interleaved 1F1B 减气泡；EP 把 896 专家分片、靠 all-to-all 做 dispatch（发 token 给专家）和 combine（收结果）；"
                 "CP 上下文并行（KCP）把 1M 序列切段、用固定大小 all-gather 同步递归状态（KDA 特有，后面单讲）；"
                 "ZeRO-1 分片优化器状态、Pipeline ZeRO-2 分片梯度还能 offload CPU。\n\n"
                 "下半是三大痛点，也是解题顺序：① EP token 负载在各 rank 间不均，热 rank 拖慢整体、动态形状还造成碎片 → MoonEP；"
                 "② 激活+梯度+优化器状态超显存 → 显存工程；③ ViT 计算方差大，暴露在关键路径会拖慢流水 → 动态 CP + 塞进气泡。"
                 "接下来逐个拆，外加 KDA 算子本身的高效实现。")

    content("叠瓦式调度：每条通信都被计算盖住", "Fig.11 调度复现",
            image="f_schedule.png", layout="big",
            blts=[("shared expert 拆 SE1/SE2 派独立 stream → 填 all-to-all 空档", 0),
                  ("reduce grad = reduce_scatter + onload + add + offload（Pipeline ZeRO-2 + CPU 梯度）", 0),
                  ("EP-DR（dispatch 重算）在反向 → memory-efficient MoE", 0)],
            takeaway="通信藏进计算的缝隙——这是后续所有优化能成立的基础。",
            note="【调度图 · 约1.5分钟】把『用计算盖住通信』落到实处（复现报告 Fig.11）。四条泳道：计算流、EP 通信、NCCL 通信、激活 offload。"
                 "三个设计意图：共享专家拆成 SE1/SE2 放独立 stream 填专家 all-to-all 的空档；那个 reduce grad 是复合操作——"
                 "reduce_scatter 后 onload CPU 分片、相加、再 offload，正是 Pipeline ZeRO-2 加 CPU 梯度；"
                 "反向里 EP-DR 是 dispatch 重算，对应后面 memory-efficient MoE。总之通信全被藏进计算缝隙里。")

    content("interleaved 1F1B + VPP：压 bubble，留一个要治的尾巴", "bubble ≈ (p−1)/(m·v)",
            image="f_1f1b.png", layout="big",
            blts=[("1F1B：不改 bubble 比例，但每 stage 只驻留约 p 份在飞激活 → 降峰值", 0),
                  ("VPP：每物理 stage 再拆 v 个虚拟块 → bubble 再除以 v，代价是通信 ×v", 0),
                  ("副作用：暖机期各 stage 驻留激活不均 → 后面用『PP 激活再平衡』治", 0)],
            takeaway="VPP 把 bubble 再除以 v，代价是暖机期激活不均（后面会治）。",
            note="【1F1B+VPP · 约1分钟】给熟悉 PP 的同事一个定量直觉。GPipe bubble 占比约 (p−1)/m；1F1B 不改比例但让每 stage 稳态只驻留约 p 份在飞激活、降峰值；"
                 "VPP 把每物理 stage 拆 v 个虚拟块，bubble 降到约 (p−1)/(m·v)，代价是通信次数 ×v。K3 层多机器多，这个收益值得。"
                 "但它留了个尾巴（右图）：暖机期各 stage 驻留激活不均，越靠前越多。这个副作用后面用『PP 激活再平衡』治，先记住。")

    content("KDA 算子①：FlashKDA —— 把慢路径榨干", "双 kernel 流水 + CHUNK=16 的三重动机",
            image="f_flashkda.png", layout="left",
            blts=[("问题：chunkwise 里 token 并行阶段被低并行度递归卡住，SM 大量闲置", 0),
                  ("双 kernel 流水：K1（token 并行：门/归一/衰减/求逆）+ K2（head 并行：跨 chunk 递归）", 0),
                  ("拆开后两阶段独立调优 + 重叠 → 端到端 ≥15% 提速", 1),
                  ("CHUNK=16：数值落进 BF16（配 lower-bounded decay）/ 16×16 逆用 Neumann / 映射 SM80 MMA 可移植", 0)],
            takeaway="呼应结构节：lower-bounded decay 让 CHUNK=16 的数值范围成立。",
            note="【FlashKDA · 约2分钟】第一个击破点：KDA 这个新算子在训练/prefill 时怎么高效——他们开源的 FlashKDA。"
                 "问题本质：chunkwise 里块内 token 并行（能喂满 SM）、块间递归（并行度低），放一个 kernel 里 token 阶段会被递归卡住、SM 闲置。"
                 "解法是拆两个 kernel 流水：K1 负责 token 并行部分（门激活、L2 归一、施加衰减、构造矩阵、16×16 求逆），"
                 "K2 负责 head 并行的跨 chunk 递归和输出；拆开后各自独立调 occupancy 和 tiling，端到端至少 15% 提速。\n\n"
                 "CHUNK=16 的三重动机（右图）：一，配合结构节的 lower-bounded decay（gmin=−5），16 token 的累积衰减正好落进 BF16、免复杂 rescale；"
                 "二，16×16 矩阵求逆用 Neumann 级数直接展开、很便宜；三，都能映射到 SM80 的 MMA 指令、多种 NVIDIA GPU 可移植。"
                 "还有精度上的抠：片上状态存 bf16 省一半 shared memory（FMA 用 fp32 就无损）。这页和结构节的 KDA③ 是一条线。")

    content("KDA 算子②：KCP —— 线性注意力的上下文并行", "转移矩阵分解 + 前缀扫描",
            image="f_kcp.png", layout="left",
            blts=[("softmax 的 CP（Ring Attention）要交换随序列增长的 KV 块；线性注意力只需传固定大小状态", 0),
                  ("但 delta rule 不能简单求和：转移矩阵 M_t 作用在入段状态上", 0),
                  ("KCP：每段分解成『累积转移 M』+『零初值局部状态 S̃』两个本地可算量", 0),
                  ("→ 一次固定大小 all-gather + 前缀扫描精确重组；通信与序列长度无关 → 线性扩展", 0)],
            takeaway="思想迁移：softmax 的 CP，用转移矩阵分解+前缀扫描搬到线性注意力。",
            note="【KCP · 约2分钟】第二个击破点：1M 序列怎么切到多设备并行。先对比通信代价：softmax 的 CP（Ring Attention）要在设备间传 KV 块、随长度增长；"
                 "朴素线性注意力很简单——每 rank 从零状态算自己这段的状态再求和。但 KDA 的 delta rule 不能这么简单求和："
                 "S_t = M_t·S_{t−1}+β_t·k_t·v_tᵀ，其中 M_t=(I−β_t·k_t·k_tᵀ)·Diag(α_t) 是作用在入段状态上的转移矩阵——"
                 "这一段对状态的影响依赖进入这段时的状态，不像纯累加那样与入段无关，所以不能直接求和。\n\n"
                 "KCP 的解法：把每段效应分解成两个都能只用本段 token 算的量——累积转移矩阵 M（描述这段怎么变换入段状态）和从零出发的局部状态 S̃，"
                 "任意入段状态进来，出段状态 = S̃ + M·入段状态。这些 rank 级更新满足结合律，于是一次固定大小 all-gather 交换各段的 (M, S̃)、"
                 "再做前缀扫描就能精确重组。通信量与序列长度无关 → 线性扩展。这是把 softmax 的 CP 思路迁到线性注意力的漂亮例子。")

    content("MoonEP①：让每个 rank 恰好算 S×K 个 token", "常规 EP 三痛，其实同源",
            image="f_moonep_balance.png", layout="left",
            blts=[("常规 EP 三痛：热 rank 决定迭代时间；动态形状→显存碎片直至 OOM；每层 host-device 同步取形状", 0),
                  ("三点同源：都是『token→专家映射动态且不均』的后果", 0),
                  ("MoonEP 用动态冗余专家做到完美均衡：热专家临时复制到别的 rank、分流 token", 0),
                  ("前向在线规划冗余专家并预取权重；反向把冗余梯度 reduce 回 home rank", 1)],
            takeaway="不去缓解不均，而是直接消灭它——后面所有红利从这来。",
            note="【MoonEP① · 约2分钟 · 已开源 · 训练 infra 核心】常规 EP（如 DeepEP）三痛：token 路由到各 rank 数量不均，最热的 rank 决定整层时间，"
                 "其他 rank 干等；每步每层路由结果都变、专家激活形状动态，导致显存碎片久了甚至 OOM；因为形状动态，host 每层要和 device 同步一次拿实际形状才能发 kernel，流水停顿。"
                 "关键洞见：这三点同源，都是『token 到专家映射动态且不均』的后果——所以治本是让这个映射结果对系统而言变成静态且完全均衡。\n\n"
                 "MoonEP 核心是『动态冗余专家』，目标让每 rank 恰好处理 S×K 个 token（S=每 rank 序列长度、K=每 token 选专家数），"
                 "所有 rank 计算量完全一样。前向根据当前 micro-batch 路由结果在线规划——某专家太热就在别的 rank 放冗余副本、分流超出的 token、并预取权重；"
                 "反向冗余副本的梯度再 reduce 回它真正的 home rank。看底部图（H20, EP=8）：MoonEP 随不均几乎平坦，DeepEP 稳步恶化。")

    content("MoonEP②：为什么『完美均衡』是可证明的", "E/R 冗余专家上界，区别于 ECHO/UltraEP 的根本",
            image="f_moonep_proof.png", layout="left",
            blts=[("定理：任意路由下，每 rank 只需 ≤ E/R 个冗余专家就一定能均衡（E=专家数, R=EP 并行度）", 0),
                  ("证明：反复用超载 rank 填满欠载 rank，每 rank 至多被填一次 → 远端 token 只来自单一 rank → 至多 E/R 个专家", 1),
                  ("下界 ≈ E/R → 界基本紧；预留 E/R 槽位后规划永远可行、训练永不中断", 0),
                  ("对比 ECHO/UltraEP：固定冗余数或 token cap，无解时被迫停训 + 手调", 1)],
            takeaway="把『确定性』当一等目标：永不中断、静态形状、零拷贝。",
            note="【MoonEP② · 约2分钟 · 数学在备注】MoonEP 最漂亮的地方是理论保证。设 E 是专家总数、R 是 EP 并行度。"
                 "他们证明：对任意路由输出，每个 rank 最多只需 E/R 个冗余专家，就一定存在完美均衡方案。\n\n"
                 "证明思路：从『每 rank 只放本地专家』开始，分成超载/欠载两类；反复挑一个超载 rank 去把某个欠载 rank 填到恰好 S×K。"
                 "每次填充让一个欠载 rank 达到平衡且此后不变，所以最多 R−1 次结束；且每 rank 最多被填一次，意味着它的远端 token 只来自单一另一个 rank，"
                 "那个 rank 上至多 E/R 个本地专家，所以冗余专家数 ≤ E/R。还证了下界约等于 E/R，界基本紧。\n\n"
                 "实际意义：只要预留 E/R 个槽位，规划永远有可行解、训练永不因『找不到均衡方案』中断。"
                 "对比 ECHO/UltraEP——预设固定冗余数或 per-rank token 上限，一旦某步路由太极端没有可行方案训练就得停，上限还要手工调、仍残留不均。"
                 "这就是 MoonEP 的工程哲学：把确定性当一等目标。")

    content("MoonEP③：完美均衡换来的两个红利", "零拷贝 + 静态形状",
            image="f_moonep_buffer.png", layout="right",
            blts=[("零拷贝：规划 kernel 预算每 token 目的地 → 直接写到远端最终槽位，buffer 视图交给计算", 0),
                  ("DeepEP 最坏要 S×K×R 的 buffer；MoonEP 完美均衡恒定 S×K，与倾斜无关", 1),
                  ("静态形状：每层 MoE 形状静态已知 → 消除 host 同步、降 kernel 启动开销", 0),
                  ("rank 内仍偏斜 → workload-aware GEMM 调度（代价模型+离线 autotune）+ 共享专家单独 stream", 1)],
            takeaway="完美均衡是因，零拷贝/静态形状/无碎片都是果。",
            note="【MoonEP③ · 约1.5分钟】完美均衡带来两个很值钱的副产品。零拷贝：因为规划 kernel 已算好每个 token 该去哪，"
                 "token 可以直接写到远端『按专家分组』的最终位置，通信 buffer 的视图直接交给计算，省掉『通信缓冲区→用户缓冲区』那次拷贝（通常是 epilogue 主开销）。"
                 "看右图：DeepEP 要支持零拷贝，最坏情况（所有 rank 把 token 发给同一 rank）得预留 S×K×R 的 buffer；MoonEP 恒定 S×K，与倾斜无关。\n\n"
                 "静态形状：每层 MoE 计算形状静态已知，就不用每层把计数读回 host 决定 kernel 形状了，消除同步、降启动开销。"
                 "不过 rank 内各专家之间还是偏斜的，用 workload-aware 的 GEMM 调度治（分析型代价模型+离线 autotune），共享专家 GEMM 单独放 stream overlap。"
                 "所以你看——均衡不只是提吞吐，它顺带把显存碎片也消掉了（动态形状本来是碎片根源）。")

    content("显存工程：把『省显存』做成可组合的存储策略", "统一激活管理器 + 账本",
            image=["f_act_manager.png", "f_mem_ledger.png"], layout="two",
            blts=[("每个为反向保存的张量挂『可插拔存储后端』：重计算 / FP8 块量化 / offload / 远程 offload，张量粒度自由组合、单内存池避免碎片", 0),
                  ("MoE：梯度数学变换免存前向输出 + 反向重算 dispatch（通信与反向计算重叠）；Block AttnRes：整体 checkpoint → 显存足迹=标准残差", 0)],
            takeaway="没有银弹——多种手段各省一点，合起来才把 2.8T 塞进预算。",
            note="【显存工程 · 约2分钟】第四个击破点：2.8T 怎么装进显存。核心是一个优雅抽象——统一激活管理器（左图）："
                 "每个需要为反向保存的张量都挂一个『可插拔存储后端』，也就是『怎么省这块显存』不再写死在模型代码里，"
                 "而是变成张量粒度上可声明的『存储策略』、和模型逻辑解耦。策略有四种：重计算、FP8 块量化、offload（到 CPU 或别的 PP rank）、"
                 "以及组合（同一张量叠加多种）；所有 GPU 内存在主计算流上从单一内存池分配，避免多流碎片；激活按层预取回来和计算重叠。\n\n"
                 "右图是账本，各手段各省一点。举几个实际用法：多数激活 FP8 块量化+offload、逐元素算子直接重计算；"
                 "MoE 借 SonicMoE 思路把某个梯度数学变换成只依赖中间激活和上游梯度、免存前向输出，group GEMM 只存 dispatch 输入、反向重算 dispatch（重算通信正好和反向计算重叠）；"
                 "Block AttnRes 整个用 checkpoint 包起来，于是每层为反向保存的激活和标准残差架构完全一样——等于白拿了 AttnRes 不额外占显存。")

    content("显存工程（续）：三个系统级手段", "把压力从 HBM 挪到 CPU / 别的卡 / P2P",
            image="f_p2p_muon.png", layout="left",
            blts=[("PP 激活再平衡：1F1B 暖机各 stage 激活不均 → Mooncake 远程 offload 到别的 PP rank 削峰", 0),
                  ("Pipeline ZeRO-2 + CPU 梯度：梯度分片下沉 CPU，GPU 只留 double buffer", 0),
                  ("P2P Muon：正交化需完整矩阵；改为只 P2P 拉『自己负责更新的』分片", 0),
                  ("消除全参 all-gather 的显存+通信双爆炸，buffer 从 O(P) 降到 O(P/DP)，按 model-chunk 流水隐藏", 1)],
            takeaway="削峰而非降总量——显存瓶颈是『单卡峰值』，不是『总量』。",
            note="【显存续 · 约2分钟】三个系统级手段，主题都是『把压力从 HBM 挪走』。PP 激活再平衡：回收 1F1B 那个尾巴——"
                 "暖机期越靠前 stage 激活越多、越靠后越少，峰值不均易 OOM，用 Mooncake Transfer Engine 把激活远程 offload 到别的 PP rank 内存、削平峰值。"
                 "关键认知：显存瓶颈是『单卡峰值』不是『总量』，把峰值从热卡挪到冷卡等效于给你更大空间。\n\n"
                 "Pipeline ZeRO-2 把梯度分片后下沉 CPU，GPU 只留 double buffer。P2P Muon（右图）——Muon 的 Newton–Schulz 正交化是整矩阵多项式，"
                 "必须完整矩阵，但优化器把参数分片了；朴素做法全量 all-gather，显存和通信双爆炸。K3 改成每卡只用点对点通信拉回『自己负责更新的』参数分片，"
                 "临时 buffer 从 O(P) 降到 O(P/DP)、消除全参 buffer，再按 model-chunk 粒度流水隐藏。这个『只有要更新该参数的卡才需要它的完整矩阵、"
                 "所以把广播降级成定向』的思路，对我们做分布式通信很有借鉴。")

    content("多模态编码器：把 ViT 藏进流水线气泡", "动态 CP + bubble filling",
            image="f_bubble_fill.png", layout="left",
            blts=[("痛点③：大图/长视频让 ViT 计算量方差大，放关键路径会拖慢流水、造成设备间负载不均", 0),
                  ("动态 CP：单张大图沿 patch 维切到多设备，跨 CP rank gather-KV 算注意力", 0),
                  ("CP 组再分 sub-group，多张大图负载均衡分散，通信占比不随规模膨胀", 1),
                  ("塞进 PP 气泡：首批 micro-batch 的 ViT 前向同步先跑，其余 ViT 前/反向调度进气泡", 0),
                  ("→ ViT 有效开销基本被隐藏（承 K2.5 DEP，思路同 Optimus）", 1)],
            takeaway="原生多模态里视觉不是附加项，要和文本在系统层深度融合。",
            note="【多模态 encoder · 约1.5分钟】第三个痛点：ViT 计算方差。背景：K3 原生多模态，文本和视觉在同一骨干联合训练，"
                 "但大图/长视频的 ViT 计算量差异巨大，直接放关键路径会拖慢流水、造成设备间负载不均（有的卡啃大图、有的卡闲着）。\n\n"
                 "两招：第一，动态上下文并行——单张大图沿 patch 维切到多设备，注意力通过跨 CP rank 的 gather-KV 算；"
                 "再把一个 CP 组分成几个 sub-group，把多张大图负载均衡分散，通信占比不随规模膨胀。"
                 "第二，把 ViT 计算塞进流水线气泡——interleaved 1F1B 下最前几个 micro-batch 的文本前向排在最开始、最后几个的反向排在最末尾，中间有气泡；"
                 "把首批 micro-batch 的 ViT 前向同步先跑掉，其余 ViT 前向和反向调度进气泡。结果 ViT 有效开销基本被隐藏、视觉几乎白算。"
                 "承接 K2.5 的 Decoupled Encoder Process，思路和 Optimus 一致。")

    content("对照：K3 vs DeepSeek 的训练 Infra 气质", "两种工程审美",
            table=[["主题", "DeepSeek（V3/V4 公开）", "Kimi K3"],
                   ["EP 通信", "DeepEP：高效 all-to-all，容忍不均", "MoonEP：动态冗余做到完美均衡 + 静态形状"],
                   ["负载均衡", "EPLB（部署期）+ sign 偏置", "训练期每步在线规划 + Quantile Balancing"],
                   ["低精度", "FP8 训练（块量化）", "FP8 存激活 + 后训练 MXFP4/8 QAT"],
                   ["流水线", "DualPipe 双向流水降 bubble", "1F1B+VPP，把 ViT/offload/EP 通信塞进 bubble"],
                   ["长上下文", "NSA/DSA/CSA 稀疏 kernel", "FlashKDA + KCP（线性注意力专属）"]],
            layout="table",
            blts=[("共同点：架构为系统而生（MLA/KDA 都是为 cache 与通信设计的注意力）", 0),
                  ("差异：K3 把训练期确定性当一等目标；DeepSeek 更极致做通信-流水线重叠", 0)],
            takeaway="没有绝对优劣——确定性 vs 极致重叠，两种审美。",
            note="【K3 vs DeepSeek Infra · 约1.5分钟 · 本节收口】逐行快过：EP 通信——DeepSeek 的 DeepEP 是『高效 all-to-all 但容忍不均』，"
                 "K3 的 MoonEP 是『动态冗余做到完美均衡还顺带静态形状』；负载均衡——DeepSeek 主要在部署期用 EPLB 再平衡+训练期 sign 偏置，"
                 "K3 是训练期每步在线规划+Quantile Balancing；低精度——DeepSeek FP8 训练，K3 训练期 FP8 存激活+后训练 MXFP4/8 QAT；"
                 "流水线——DeepSeek 有 DualPipe 双向流水极致压气泡，K3 用 1F1B+VPP 把 ViT/offload/EP 通信塞进气泡；"
                 "长上下文——DeepSeek 是 NSA/DSA/CSA 稀疏 attention kernel，K3 是 FlashKDA+KCP 线性注意力专属。\n\n"
                 "两点总结：共同点是『架构为系统而生』——MLA 和 KDA 本质都是为 cache 和通信设计的注意力，不是纯建模视角；"
                 "差异在气质——K3 强调训练期确定性（完美均衡、静态形状、永不中断），DeepSeek 强调把通信和流水线重叠做到极致。没有绝对优劣。")

    # ============ 05 · RL ============
    slide_section("05", "强化学习（简化）",
                  "SFT → 分域分强度 RL → 多教师蒸馏；重点看『范式』而非细节",
                  "【过场 · 约30秒】强化学习简化处理，只讲骨架和几个关键设计，细节留给文档。三阶段：SFT 冷启动 → 分域、分推理强度做 RL 得一批专家 → "
                  "多教师在线蒸馏合回一个模型。记住这个『分而治之再合并』的范式就够了，它已是 2026 前沿共识。")

    content("后训练三阶段：分而治之，再合并", "9 个专家 → 一个统一模型",
            image="f_rl_experts.png", layout="left",
            blts=[("SFT：合成复杂 agentic 轨迹 + XTML 模板统一序列化；QAT 从此开始（MXFP4/8）", 0),
                  ("RL：3 域 × 3 reasoning effort = 9 个专家", 0),
                  ("RL FLOPs↑ → 工具调用步数与能力同步增长（长程自主执行是学出来的）", 1),
                  ("MOPD：多教师 on-policy 蒸馏合并 —— 与 DeepSeek-V4 / MiMo-V2 范式收敛", 0)],
            takeaway="不给每个任务单独训，而是『域×强度』九宫格再蒸回一个。",
            note="【后训练流程 · 约2分钟】SFT 阶段：用前代 Kimi 域专家合成复杂 agentic 轨迹，加多级校验和人在环标注；"
                 "所有数据用 XTML 模板序列化——为 Agent 时代重设计的 chat template，用三个保留 token 表示结构边界、消除 tokenization 歧义、方便约束解码，"
                 "而且低对齐税、轻量 SFT 后就能上 RL。QAT 从 SFT 就开始（MoE 权重 MXFP4、激活 MXFP8），一直贯穿到 RL、训推同量化。\n\n"
                 "RL 阶段：不给每个任务单独训，而是分三大域（通用、通用 Agent、编码 Agent）× 三档推理强度（low/high/max）得 9 个专家（图中 3×3 网格）。"
                 "一个重要现象：RL FLOPs 增加时，工具调用步数和综合能力同步上涨——『更长的自主执行』本身是 RL 学出来的。"
                 "MOPD 阶段用多教师在线蒸馏把 9 个专家合回一个统一模型。这个范式和 DeepSeek-V4、小米 MiMo-V2 收敛，已成前沿共识。")

    content("RL 的关键设计与环境（快速过）", "四个关键设计 + 六类环境",
            table=[["统一白盒 harness", "知识图谱任务合成", "Kernel 优化(+反 hacking)"],
                   ["个人助理 mock 应用", "AET 自主执行", "WebDev 多 scaffold"]],
            layout="table",
            blts=[("Partial rollout：完成 λ 比例即优化，未完轨迹入队下轮续跑 → 长尾不阻塞；per-token 正则容忍极端 off-policy", 0),
                  ("Reasoning-effort 预算：超 τ·b₀(x) 的轨迹 reward=−1；τ 退火得 max/high/low", 0),
                  ("Agentic GRM（不可验证任务）：锦标赛两两比较 + rubric 协议 + 长度阈值防刷分", 0),
                  ("训练+评测全程创建 5,100 万+ 沙箱 / 150 万+ 镜像（AgentENV microVM）——环境工程=核心壁垒", 0)],
            takeaway="环境工程（沙箱/harness/任务合成）是新的核心壁垒。",
            note="【RL 关键设计 · 约2分钟 · 快速过】四条关键设计。Partial rollout：长程任务总有个别轨迹特别慢，每轮采样一批、完成到 λ 比例就开始优化、"
                 "没完成的入队下轮接着跑（靠沙箱恢复状态），长尾不阻塞；副作用是长轨迹跨多轮变成极端 off-policy，用 per-token 正则把更新约束在局部邻域、容忍过期数据。"
                 "Reasoning-effort 预算：给每题估初始 token 预算 b₀(x)，超 τ·b₀(x) 直接给 −1，τ 从大到小退火得 max/high/low 三档；"
                 "通用任务只数思考 token、agentic 数全部输出（含工具参数），防止把话藏进工具参数绕过限制。"
                 "Agentic GRM：不可验证任务用生成式奖励模型做锦标赛式两两比较，强制 judge 走『读产物→生成 rubric→逐项打分→记分板』协议，超长直接判负防刷分。\n\n"
                 "下半六类环境只念名字、点一句『环境工程是核心壁垒』：统一白盒 harness（可实例化 Kimi Code/Claude Code/Codex，防过拟合单一 harness）、"
                 "知识图谱引导任务合成、kernel 优化（含反 reward-hacking）、个人助理 mock、AET 自主执行、WebDev 多 scaffold。"
                 "训练+评测全程创建了五千多万个沙箱、一百五十多万个镜像，用的是 AgentENV microVM。")

    content("蒸馏合并、量化与投机解码", "让训练目标对齐部署现实",
            layout="none",
            blts=[("MOPD：学生在自己分布上采样、教师逐 token 打分（dense reward）→ 避免 SFT 蒸馏的 exposure bias", 0),
                  ("reward = 教师/学生对数概率比，clip 到 ±R_max；无缝接入 RL 框架、复用 partial rollout", 1),
                  ("QAT 贯穿 SFT+RL：MoE 权重 MXFP4 / 激活 MXFP8，非专家部分高精度；rollout 与训练同量化 → 训推一致", 0),
                  ("投机解码：复用预训练 MTP 层做 EAGLE-3 draft；LK loss 直接优化接受率而非 KL 代理", 0),
                  ("贯穿主线：QAT 对齐量化 · LK loss 对齐接受率 · on-policy 对齐推理分布", 0)],
            takeaway="训练目标尽量对齐部署现实——这是很多 RL 翻车的隐形来源。",
            note="【蒸馏/量化/投机 · 约1.5分钟】后训练收尾三件事，都体现同一主线：让训练目标对齐部署现实。"
                 "MOPD（多教师在线蒸馏）：学生在自己分布上采样（on-policy），教师对学生生成的每个 token 打分，避免 SFT 式离线蒸馏的 exposure bias。"
                 "公式（备注）：per-token reward = clip(stop_grad(log[π_teacher/π_student]), −R_max, R_max)，是 dense 的、能无缝接入现有 RL 框架、复用 partial rollout。\n\n"
                 "QAT 从 SFT 到 RL 全程，MoE 权重 MXFP4、激活 MXFP8，非专家部分高精度；关键是 RL 阶段 rollout 和训练用同一套量化、做到训推一致——"
                 "这是很多 RL 训练翻车的隐形来源。投机解码：K3 预训练就带一个 MTP 层，结构正好是单层 decoder block，"
                 "冻结主模型微调成 EAGLE-3 draft；用 LK loss 直接优化接受率而非 KL 代理，因为接受率才是投机解码真正的加速指标。")

    # ============ 06 ============
    slide_section("06", "评测、成本与总结",
                  "开源新前沿的证据，以及可以带走的方法论",
                  "【过场 · 约30秒】最后一节给结果与收束。三块：评测（K3 在开源什么位置、强在哪弱在哪）、成本效率（把架构选择兑现成商业价值、和开场呼应）、"
                  "几个能力案例（模型反哺 infra 的飞轮），最后用五条可带走的方法论收尾。时间紧的话评测和案例各半分钟带过，把重点留给最后那页。")

    content("评测与成本效率", "开源第一梯队 + 成本-质量前沿",
            image="f_cost.png", layout="left",
            blts=[("整体：Fable 5 ≳ GPT-5.6 Sol > Kimi K3 > Opus 4.8 / GPT-5.5 / GLM-5.2；研究级推理仍有差距", 0),
                  ("亮点：BrowseComp 91.2（第一）· ProgramBench 77.8（第一）· SWE-Marathon 42.0（领先 Fable 5 七点）", 0),
                  ("第三方：AA 智能指数 #4/580；WebDev Arena 开源首次登顶（Elo 1678）", 0),
                  ("成本：BrowseComp 最高分且仅 $2.03/任务 = GPT-5.6 Sol 的一半、Claude max 档的十分之一量级", 0)],
            takeaway="架构效率最终兑现为价格优势——呼应开场。",
            note="【评测与成本 · 约1分钟】整体排名 K3 稳居开源第一梯队，仅次于 Claude Fable 5 和 GPT-5.6 Sol，领先其余所有；短板在研究级推理（HLE、CritPt）。"
                 "亮点是 BrowseComp、ProgramBench 拿第一，SWE-Marathon（GPU kernel 向）领先 Fable 5 七个点。第三方里 Artificial Analysis 智能指数 580 模型排第 4，"
                 "WebDev Arena 是开源模型首次登顶。\n\n"
                 "更值得讲的是成本（右图，横轴每任务成本、纵轴分数，越左上越好）：K3 基本落在成本-质量前沿上。以 BrowseComp 为例，拿最高分同时只要 2.03 美元一个任务，"
                 "是 GPT-5.6 Sol 的一半、Claude max 档的十分之一量级。这不是偶然——正是前面所有『为系统让路』的架构选择（KDA 固定状态、MLA 低秩、混合注意力）在服务成本上的兑现。"
                 "可以呼应开场：架构效率最终变成商业上的价格优势。")

    content("案例：模型开始反哺自己的 Infra", "能力-Infra 飞轮",
            image="f_kernel_case.png", layout="left",
            blts=[("GPU kernel 优化（24h/任务）：AttnRes 283.6→114.4ms；KDA −73.6%；追平 Fable 5", 0),
                  ("早期 K3 checkpoint 已承担团队内多数 kernel 优化工作", 1),
                  ("MiniTriton：自研 DSL→MLIR→PTX 编译器；L20 上 matmul 达 cuBLAS ~90%", 0),
                  ("nano-kpu：48h 自主跑通开源 EDA，4mm²/100MHz/8700 tok·s⁻¹ 芯片原型", 0),
                  ("安全：Linux kernel 发现 16 个未知漏洞（含远程堆越界写、Dirty-COW 类提权）", 0)],
            takeaway="用 RL 环境练 kernel → 模型替团队写 kernel → 支撑下一代训练。",
            note="【案例 · 约1分钟】主题是『模型开始反哺自己的基础设施』。GPU kernel 优化：给每个模型 24 小时独立优化 kernel，"
                 "K3 把 AttnRes kernel 从 283.6ms 降到 114.4ms、KDA 降 73.6%、整体追平带 fallback 的 Claude Fable 5；"
                 "报告还提到一个早期 K3 checkpoint 就已经承担团队内部大部分 kernel 优化工作了。"
                 "MiniTriton：K3 自己写了个类 Triton 编译器，DSL 前端→MLIR→PTX 全链路，L20 上 matmul 达 cuBLAS 约 90%。"
                 "nano-kpu：48 小时自主跑通开源 EDA 设计出推理芯片原型，4mm²、100MHz 时序收敛、RTL 仿真解码 8700 tok/s。"
                 "安全评测在 Linux kernel 发现 16 个未知漏洞（含可远程触发的堆越界写和 Dirty-COW 类提权）。"
                 "串成一句话：用 RL 环境训练 kernel 能力 → 模型替团队写 kernel → 这些 kernel 又支撑下一代训练——能力和基础设施的飞轮已经转起来了。")

    content("总结：可以带走的五条", "一页收束全场",
            table=[["#", "带走", "一句话"],
                   ["1", "为什么这样设计", "KDA/AttnRes/LatentMoE 各解一维信息流瓶颈；混合注意力解耦『成本 vs 检索』"],
                   ["2", "算法-系统协同", "lower-bounded decay 换 Tensor Core、NoPE 免调参、QAT 对齐量化——每层为下一层让路"],
                   ["3", "确定性即性能", "MoonEP 完美均衡+静态形状+永不中断；把随机性与长尾挤出训练"],
                   ["4", "方法论可迁移", "比较技术前先各自调最优（cosine vs WSD）；存储策略抽象；QB 直方图均衡"],
                   ["5", "能力-Infra 飞轮", "模型优化 kernel → 支撑下一代训练；环境/沙箱工程成核心壁垒"]],
            layout="table",
            blts=[("一句话：K3 用 2.8T 证明了『线性混合路线』能撑起 frontier 模型 + 1M 上下文 + 可负担的服务成本", 0)],
            takeaway="最好的系统优化，往往始于一个小的算法改动。",
            note="【总结五条 · 约1分钟收尾】用五条把全场收起来："
                 "一，为什么这样设计——KDA、AttnRes、LatentMoE 各解序列/深度/宽度一个瓶颈，混合注意力把成本和精确检索解耦。"
                 "二，算法-系统协同，贯穿全场——lower-bounded decay 为换 Tensor Core、NoPE 为免调参、QAT 为对齐量化、LK loss 为对齐接受率，每层抽象为下一层让路。"
                 "三，确定性即性能——MoonEP 的完美均衡、静态形状、永不中断，本质是把随机性和长尾从训练里挤出去。"
                 "四，方法论可迁移——最想让大家带走的是『比较任意训练技术前先各自调到最优』（cosine vs WSD 那个坑），还有统一激活管理器的存储策略抽象、QB 的直方图均衡。"
                 "五，能力-Infra 飞轮——模型已经能优化 kernel、反哺训练，环境和沙箱工程成为新壁垒。"
                 "最后一句：K3 用 2.8T 证明了线性混合路线能撑起 frontier 模型 + 1M 上下文 + 可负担的服务成本。开放提问。")

    # Q&A
    s = prs.slides.add_slide(BLANK); bg(s, NAVY_D); hero_full(s, dim=40)
    _, tf = textbox(s, 0.9, 2.9, 11.5, 2.4)
    add_para(tf, "Q & A · 谢谢！", size=44, color=WHITE, bold=True, first=True, space_after=10)
    add_para(tf, "常见问题预案：KDA vs Mamba/GLA（delta rule 纠错 + 通道级门）· 为何不用纯稀疏 attention · "
             "MoonEP 与 DeepEP/EPLB 差异 · NoPE 外推 · 896 选 16 如何训得起", size=14,
             color=RGBColor(0xCB, 0xD5, 0xE1))
    add_para(tf, "细节与公式详解见 DESIGN_NOTES.md / PRETRAIN_INFRA_DEEPDIVE.md", size=13,
             color=RGBColor(0x5E, 0xEA, 0xD4))
    _idx[0] += 1
    notes(s, "【Q&A】谢谢大家。常见问题预案见页面。更细的推导和对比在配套文档里：DESIGN_NOTES 是全量思路，"
             "PRETRAIN_INFRA_DEEPDIVE 是 Infra 深挖。")

    out = os.path.join(HERE, "kimi-k3-tech-share-v3.pptx")
    prs.save(out)
    print("saved:", out, "| slides:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    build()

